#-*-coding:gb2312-*-

import tkinter
import os
import docx
import re
import threading
import datetime
import shutil
import win32com
import xlsxwriter as xw

from win32com.client import Dispatch
from tkinter import *
from tkinter import ttk

full_docx = []
full_comments = []
txt_line = []
comments_list = []
true_comments_list =[]
false_comments_list =[]

class Comments(): #批注类
    def __init__(self, filename = "", filepath = "", page = "", line = "", txt = "", comments = "", date = "", author = "", done = "False"):
        self.filename = filename
        self.filepath = filepath
        self.page = page
        self.line = line
        self.txt = txt
        self.comments = comments
        self.date = date
        self.author = author
        self.done = done
    pass
    
    def __str__(self):
        ret  = "++++++++++++++++++++++++++++\n"
        ret += "filename: " + self.filename + "\n"
        ret += "filepath: " + self.filepath + "\n"
        ret += "page: " + self.page + "\n"
        ret += "line: " + self.line + "\n"
        ret += "txt: " + self.txt + "\n"
        ret += "comments: " + self.comments + "\n"
        ret += "date: " + self.date + "\n"
        ret += "author: " + self.author + "\n"
        ret += "done: " + self.done + "\n"
        ret += "----------------------------\n"
        return ret
    pass
    
    def add_filename(self, filename):
        self.filename = filename
    def add_filepath(self, filepath):
        self.filepath = filepath
    def add_page(self, page):
        self.page = str(page)
    def add_line(self, line):
        self.line = str(line)
    def add_txt(self, txt):
        self.txt += txt
    def add_comments(self, comments):
        self.comments += comments
    def add_date(self, date):
        self.date = str(date)
    def add_author(self, author):
        self.author = author
    def add_done(self, done):
        self.done = str(done)
pass

def txt_merge(txt): #把txt文件merge到目前的list里
    with open (txt, "r") as hd:
        from_line = hd.readlines()
        txt_line.extend(from_line)
pass

def get_process_files(root_dir): #递归得到所有word文档
    """return all files in directory"""
    cur_dir=os.path.abspath(root_dir)
    file_list=os.listdir(cur_dir)
    process_list=[]
    dir_extra_list = []

    for file in file_list:
        fullfile=cur_dir+"\\"+file
        #print(fullfile)
        if os.path.isfile(fullfile) and fullfile.endswith(".docx"):
            process_list.append(fullfile)
            #print("add " + fullfile)
        elif os.path.isdir(fullfile):
            dir_extra_list.extend(get_process_files(fullfile))

    if len(dir_extra_list)!=0:
        for x in dir_extra_list:
            process_list.append(x)

    return process_list
pass

def update_content(url): #打开word，执行宏命令
    ret = ""
    docApp = win32com.client.DispatchEx('Word.Application')
    try:
        doc = docApp.Documents.Open(url)
        #print("打开文档")
        doc.Application.Run('exportWordComments_Click')
        #print("执行宏完成")
        doc.Save()
        ret = url + " 处理成功"
    except Exception as e:
        print(e + ", 执行失败")
        ret = url + " 处理失败"
    docApp.Quit()
    return ret
pass

def proc_txt(list): #处理生成的txt临时文件
    for txt in list:
        if os.path.exists(txt):
            txt_merge(txt)
            print("merge完成： " + txt)
            os.remove(txt)
            print("remove完成：" + txt)
    to = log_path + "\\" + tag + ".txt"
    with open(to, "w") as hd:
        for line in txt_line:
            hd.write(line)
pass

def log_info_get(): #提取log里的信息
    log = txt_name #"D:\MyWork\python\get_comments_v2\log\Date_20220602_173646.txt"   
    filename = ""
    filepath = ""
    page = ""
    lines = ""
    txt = ""
    comments = ""
    date = ""
    author = ""
    done = ""
    
    with open (log, "r") as handle:
        hd = handle.readlines()
        for line in hd:          
            re1 = re.search(r"^\=+$", line)
            re2 = re.search(r"^\s*$", line)
            nouse = re1 or re2
            use   = not nouse
            if use:
                pre_flag = "GET_TXT"
                re0 = re.match(r"GET_FILENAME: (.*)", line)
                re1 = re.match(r"GET_FILEPATH: (.*)", line)
                re2 = re.match(r"GET_PAGE: (.*)", line)
                re3 = re.match(r"GET_LINE: (.*)", line)
                re4 = re.match(r"GET_TXT: (.*)", line)
                re5 = re.match(r"GET_COMMENTS: (.*)", line)
                re6 = re.match(r"GET_DATE: (.*)", line)
                re7 = re.match(r"GET_AUTHOR: (.*)", line)
                re8 = re.match(r"GET_DONE: (.*)", line)
                if re0:
                    filename = str(re0.group(1))
                elif re1:
                    filepath = str(re1.group(1))
                elif re2:
                    page = str(re2.group(1))
                elif re3:
                    lines = str(re3.group(1))                  
                elif re4:
                    txt = str(re4.group(1))
                elif re5:
                    comments = str(re5.group(1))                
                elif re6:
                    date = str(re6.group(1))
                elif re7:
                    author = str(re7.group(1))                
                elif re8:
                    done = str(re8.group(1))
                    comment = Comments(filename, filepath, page, lines, txt, comments, date, author, done)
                    #print(comment)
                    comments_list.append(comment)
                else:   
                    if pre_flag == "GET_TXT":
                        txt += " " + line.strip() 
                    elif pre_flag == "GET_COMMENTS":
                        comments += " " + line.strip()
pass

def gen_excel_mode0(): #生成excel
    output = excel_name #"D:\MyWork\python\get_comments_v2\log\Date_20220602_173646.xlsx"
    workbook = xw.Workbook(output)
    worksheet1 = workbook.add_worksheet("所有批注")
    worksheet1.activate()
    title = ['文件名', 
             '批注内容', 
             '原文', 
             '是否解决', 
             '批注者', 
             '页', '行', 
             '日期', 
             '文件路径']
    bold = workbook.add_format({
        'bold':  True,  # 字体加粗
        'border': 1,  # 单元格边框宽度
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'fg_color': '#F4B084',  # 单元格背景颜色
        'text_wrap': True,  # 是否自动换行
    })
    worksheet1.write_row('A1', title, bold)
    
    bold = workbook.add_format({
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'text_wrap': True,  # 是否自动换行
    })   
    i = 2
    for j in range(len(comments_list)):
        comments = comments_list[j]
        insertData = [comments.filename, 
                      comments.comments,
                      comments.txt,
                      comments.done,
                      comments.author,
                      comments.page,
                      comments.line,
                      comments.date,
                      comments.filepath
                      ]
        row = 'A' + str(i)
        worksheet1.write_row(row, insertData, bold)
        i += 1
    workbook.close()
pass

def gen_excel_mode1(): #生成excel，只提取未解决的批注
    output = excel_name #"D:\MyWork\python\get_comments_v2\log\Date_20220602_173646.xlsx"
    workbook = xw.Workbook(output)
    worksheet1 = workbook.add_worksheet("未解决批注")
    worksheet1.activate()
    title = ['文件名', 
             '批注内容', 
             '原文', 
             '是否解决', 
             '批注者', 
             '页', '行', 
             '日期', 
             '文件路径']
    bold = workbook.add_format({
        'bold':  True,  # 字体加粗
        'border': 1,  # 单元格边框宽度
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'fg_color': '#F4B084',  # 单元格背景颜色
        'text_wrap': True,  # 是否自动换行
    })
    worksheet1.write_row('A1', title, bold)
    
    bold = workbook.add_format({
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'text_wrap': True,  # 是否自动换行
    })   
    i = 2
    for j in range(len(comments_list)):
        comments = comments_list[j]
        insertData = [comments.filename, 
                      comments.comments,
                      comments.txt,
                      comments.done,
                      comments.author,
                      comments.page,
                      comments.line,
                      comments.date,
                      comments.filepath
                      ]
        row = 'A' + str(i)
        if comments.done == "False":
            worksheet1.write_row(row, insertData, bold)
            i += 1
    workbook.close()
pass

def gen_excel_mode2(): #生成excel，分两页提取
    output = excel_name #"D:\MyWork\python\get_comments_v2\log\Date_20220602_173646.xlsx"
    workbook = xw.Workbook(output)
    worksheet1 = workbook.add_worksheet("未解决批注")
    worksheet1.activate()
    title = ['文件名', 
             '批注内容', 
             '原文', 
             '是否解决', 
             '批注者', 
             '页', '行', 
             '日期', 
             '文件路径']
    bold = workbook.add_format({
        'bold':  True,  # 字体加粗
        'border': 1,  # 单元格边框宽度
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'fg_color': '#F4B084',  # 单元格背景颜色
        'text_wrap': True,  # 是否自动换行
    })
    worksheet1.write_row('A1', title, bold)
    
    bold = workbook.add_format({
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'text_wrap': True,  # 是否自动换行
    })   
    i = 2
    for j in range(len(comments_list)):
        comments = comments_list[j]
        insertData = [comments.filename, 
                      comments.comments,
                      comments.txt,
                      comments.done,
                      comments.author,
                      comments.page,
                      comments.line,
                      comments.date,
                      comments.filepath
                      ]
        row = 'A' + str(i)
        if comments.done == "False":
            worksheet1.write_row(row, insertData, bold)
            i += 1
    
    worksheet1 = workbook.add_worksheet("已解决批注")
    worksheet1.activate()
    title = ['文件名', 
             '批注内容', 
             '原文', 
             '是否解决', 
             '批注者', 
             '页', '行', 
             '日期', 
             '文件路径']
    bold = workbook.add_format({
        'bold':  True,  # 字体加粗
        'border': 1,  # 单元格边框宽度
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'fg_color': '#F4B084',  # 单元格背景颜色
        'text_wrap': True,  # 是否自动换行
    })
    worksheet1.write_row('A1', title, bold)
    
    bold = workbook.add_format({
        'align': 'left',  # 水平对齐方式
        'valign': 'vcenter',  # 垂直对齐方式
        'text_wrap': True,  # 是否自动换行
    })   
    i = 2
    for j in range(len(comments_list)):
        comments = comments_list[j]
        insertData = [comments.filename, 
                      comments.comments,
                      comments.txt,
                      comments.done,
                      comments.author,
                      comments.page,
                      comments.line,
                      comments.date,
                      comments.filepath
                      ]
        row = 'A' + str(i)
        if comments.done == "True":
            worksheet1.write_row(row, insertData, bold)
            i += 1
    workbook.close()
pass

def gen_excel(mode = 0):
    if mode == 0:
        gen_excel_mode0()
    elif mode == 1:
        gen_excel_mode1()
    else:
        gen_excel_mode2()
pass

def update_root(): #初始化结构
    global soft_root
    global log_path
    global tag
    global txt_name
    global excel_name
    
    full_docx = []
    full_comments = []
    txt_line = []
    comments_list = []
    true_comments_list =[]
    false_comments_list =[]
    
    if not os.path.exists(log_path):
        os.makedirs(log_path)
    now_time = datetime.datetime.now()
    tag = datetime.datetime.strftime(now_time,'Date_%Y%m%d_%H%M%S')
    txt_name = log_path + "\\" + tag + ".txt"
    excel_name = log_path + "\\" + tag + ".xlsx"
pass

def tk_main(): #图形界面配置
    global mode
    mode = 0 #批注提取方式
    
    def get_path():
        text1.delete("1.0", "end")
        from tkinter import filedialog
        tk_file_path = filedialog.askdirectory() #获得选择好的文件夹
        text1.insert(INSERT, tk_file_path)
    pass
    
    def proc_file(list):
        for file in list:
            tmp = re.search(r"(.*)\.docx", file).group(1)
            txt_path = tmp + "_comments.txt"
            full_comments.append(txt_path)
            #print(txt_path)
            text3.mark_set('here',1.0)
            text3.insert('here', update_content(file) + "\n")
        text3.mark_set('here',1.0)
        text3.insert('here', "==========================================================================\n")
        text3.mark_set('here',1.0)
        text3.insert('here', "全部文件处理完成，原始log路径：" + txt_name + "\n")
        proc_txt(full_comments)
    pass
    
    def start_check():      
        update_root()
        text3.delete("1.0", "end")
        text3.mark_set('here',1.0)
        text3.insert('here', "开始检索文件并处理，用时较长请勿退出，请在检查结束后点击 打开结果\n")
        fullpath = text1.get(1.0, "end").strip()
        full_docx = get_process_files(fullpath)
        proc_file(full_docx)
        log_info_get()
        gen_excel(mode)
        
        text3.mark_set('here',1.0)
        text3.insert('here', "EXCEL已生成：" + excel_name + "\n")
        text3.mark_set('here',1.0)
        text3.insert('here', "==========================================================================\n")
    pass
    
    def open_xlsx():
        already_open = 0
        xl_app = win32com.client.DispatchEx("Excel.Application")
        xl_app.Visible = 1
        for wb in xl_app.Workbooks:
            if(wb.Name == excel_name): #wb.Name只返回文件的名字，不包含路径
                already_open = 1
                break
        if(already_open==0):#需要新打开文件
            my_wb = xl_app.Workbooks.Open(excel_name)
    pass
    
    def thread_open_xlsx():
        t2 = threading.Thread(target=open_xlsx,args=())
        t2.start()
    pass
    
    def thread_start_check():
        t1 = threading.Thread(target=start_check,args=())
        t1.start()
    pass
    
    def log_shutil():
        text3.delete("1.0", "end")
        text3.mark_set('here',1.0)
        if os.path.exists(log_path):
            shutil.rmtree(log_path)
        if not os.path.exists(log_path):
            os.makedirs(log_path)
        text3.mark_set('here',1.0)
        text3.insert('here', log_path + "已清空\n")
    pass
    
    def choose(event): #选择框事件
        global mode
        mode = com.current()
        #print("value的值:{}".format(com.get()))
        #print("current:{}".format(com.current()))
    pass
    
    #tk window
    root = Tk()
    root.geometry("600x400")
    root.title("用起来不一般的word批注整理器 From GJM")
    
    f1 = Frame(root, height = 100, width = 400)
    f1.pack()
    button1 = Button(f1, text='选择目录', command=get_path)
    button1.pack(side = LEFT)
    
    #创建下拉菜单
    xVariable = tkinter.StringVar()
    com = ttk.Combobox(f1, textvariable=xVariable, cursor="arrow")
    com["value"] = ("提取全部批注", "只提取未解决批注", "分页签提取全部批注")
    com.current(0)
    com.bind("<<ComboboxSelected>>", choose)
    com.pack(side = RIGHT)
    
    text1   = Text(f1, height = 1, undo=True, autoseparators=False, width = 50)
    text1.pack(side = RIGHT)
    

    
    f3 = Frame(root, height = 100, width = 400)
    f3.pack()
    button2 = Button(f3, text='开始检查', command=thread_start_check)
    button2.pack(side=LEFT)
    button4 = Button(f3, text='打开结果', command=thread_open_xlsx)
    button4.pack(side=LEFT)
    button4 = Button(f3, text='清空缓存', command=log_shutil)
    button4.pack(side=LEFT)
    button3 = Button(f3, text='退出程序', command=root.quit)
    button3.pack(side=RIGHT)
    
    f4 = Frame(root, height = 50, width = 400)
    f4.pack()
    text3 = Text(f4, height = 50, undo=True, autoseparators=False)
    text3.pack(side = RIGHT)
    #tk window over
    
    root.mainloop()
pass

def main():
    global soft_root
    global log_path
    global tag
    global txt_name
    global excel_name

    soft_root = os.path.split(os.path.realpath(__file__))[0]
    log_path = soft_root + "\\log"
    
    if not os.path.exists(log_path):
        os.makedirs(log_path)
    now_time = datetime.datetime.now()
    tag = datetime.datetime.strftime(now_time,'Date_%Y%m%d_%H%M%S')
    txt_name = log_path + "\\" + tag + ".txt"
    excel_name = log_path + "\\" + tag + ".xlsx"
    
    tk_main()
pass

if __name__ == '__main__':
    main()
